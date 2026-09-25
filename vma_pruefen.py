#!/usr/bin/env python3
"""Prueft erzeugte VMA-Dateien gegen die Regeln aus .claude/skills/vma-zeiterfassung.

Nutzung:
  python3 vma_pruefen.py <Monat> <Jahr> <Ausgabeordner> --vorlage <xlsx> [--absences <csv>] [-v]
Exit-Code 0 nur, wenn alle Dateien OK sind.
"""
import calendar
import sys
from collections import Counter
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from vma_automation import (EMPLOYEES, get_brandenburg_holidays, hhmm_to_minutes,
                            read_absences_from_file)

MONATE = {1: 'Januar', 2: 'Februar', 3: 'Maerz', 4: 'April', 5: 'Mai', 6: 'Juni', 7: 'Juli',
          8: 'August', 9: 'September', 10: 'Oktober', 11: 'November', 12: 'Dezember'}


def arg(name):
    return sys.argv[sys.argv.index(name) + 1] if name in sys.argv else None


def zahl(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def minuten(v):
    return int(v) * 60 + round((v - int(v)) * 100)


def pruefe(emp, pfad, vorlage, month, year, absences, holidays, verbose):
    cfg = EMPLOYEES[emp]
    ws = openpyxl.load_workbook(pfad)['Tabelle1']
    probleme = []
    num_days = calendar.monthrange(year, month)[1]

    if ws['D3'].value != year or ws['E3'].value != month:
        probleme.append('D3/E3 falsch')
    if ws['A9'].value != cfg['fullname']:
        probleme.append('Name A9 falsch')
    if ws['H3'].value != f'für Monat {month:02d}/{year % 100:02d}':
        probleme.append('Monatstext H3 falsch')

    eingabe = {f'{get_column_letter(c)}{r}' for c in range(2, 33) for r in range(14, 19)}
    for row in vorlage.iter_rows():
        for c in row:
            if (isinstance(c.value, str) and c.value.startswith('=')
                    and c.coordinate not in eingabe and ws[c.coordinate].value != c.value):
                probleme.append(f'Formel {c.coordinate} geaendert')
    if any(c.number_format == '0,00' for r in ws.iter_rows() for c in r):
        probleme.append('Zahlenformat 0,00')

    abw = {d: art for d, art, _ in absences.get(emp, [])}
    werte, wochen = [], Counter()
    for d in range(1, num_days + 1):
        col = get_column_letter(d + 1)
        arbeit, urlaub, feiertag, krank = (ws[f'{col}{r}'].value for r in (14, 15, 16, 17))
        wtag = date(year, month, d).weekday()
        abwesend = any(zahl(x) for x in (urlaub, feiertag, krank))
        if zahl(arbeit) and abwesend:
            probleme.append(f'Tag {d}: Arbeitszeit und Abwesenheit')
        if d in holidays and wtag < 5 and not zahl(feiertag):
            probleme.append(f'Tag {d}: Feiertag fehlt')
        if zahl(feiertag) and (d not in holidays or wtag >= 5):
            probleme.append(f'Tag {d}: Feiertag falsch')
        if d in abw and not (d in holidays and wtag < 5):
            soll = krank if abw[d] == 'krank' else urlaub
            if not zahl(soll):
                probleme.append(f'Tag {d}: {abw[d]} fehlt')
        if zahl(arbeit):
            m = minuten(arbeit)
            werte.append(m)
            wochen[date(year, month, d).isocalendar()[1]] += m
            if m % 5:
                probleme.append(f'Tag {d}: kein 5-Minuten-Schritt')
            if m < cfg['short_range'][0]:
                probleme.append(f'Tag {d}: zu kurz ({m // 60},{m % 60:02d})')

    summe = sum(werte)
    if summe > hhmm_to_minutes(cfg['max_hours']):
        probleme.append('ueber Monatsmaximum')
    if wochen and max(wochen.values()) > 2400:
        probleme.append('Woche ueber 40:00')
    if werte:
        wert, n = Counter(werte).most_common(1)[0]
        if n > max(4, len(werte) * 0.3):
            probleme.append(f'Muster: {wert // 60},{wert % 60:02d} {n}x')
    vma = sum(1 for m in werte if m > 480)
    lo, hi = cfg['vma_target']
    hinweis = '' if lo <= vma <= hi else f' (Ziel {lo}-{hi})'

    print(f'{emp:8s} {summe // 60},{summe % 60:02d} / {cfg["max_hours"]:.2f}'.replace('.', ',')
          + f'  VMA {vma}{hinweis}  ' + ('OK' if not probleme else 'FEHLER: ' + '; '.join(probleme)))
    if verbose:
        print('   ', ' '.join(f'{m // 60},{m % 60:02d}' for m in werte))
    return not probleme


def main():
    if len(sys.argv) < 4 or '--vorlage' not in sys.argv:
        print(__doc__)
        sys.exit(2)
    month, year, ordner = int(sys.argv[1]), int(sys.argv[2]), Path(sys.argv[3])
    vorlage = openpyxl.load_workbook(arg('--vorlage'))['Tabelle1']
    absences = read_absences_from_file(Path(arg('--absences')), year, month) if arg('--absences') else {}
    holidays = get_brandenburg_holidays(year, month)
    alle_ok = True
    for emp in EMPLOYEES:
        pfad = ordner / f'Zeiterfassung_{MONATE[month]}_{year}_{emp}_fertig.xlsx'
        if not pfad.exists():
            continue
        alle_ok &= pruefe(emp, pfad, vorlage, month, year, absences, holidays, '-v' in sys.argv)
    sys.exit(0 if alle_ok else 1)


if __name__ == '__main__':
    main()
